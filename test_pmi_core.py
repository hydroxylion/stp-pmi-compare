# -*- coding: utf-8 -*-
"""pmi_core 回归测试。

用法::

    python test_pmi_core.py                    # 只跑纯函数单测（无需外部文件）
    python test_pmi_core.py <SFA报告.xlsx>      # 追加集成测试（真实数据端到端）

纯函数单测：乱码修复 / 数值归一 / 语义指纹 / 比对判定 / markdown 解析
集成测试  ：SFA 五表索引 + 三条 ID 关联路径 + 三层指标
"""
import io, json, os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import pmi_core as C

DEFAULT_XLSX = r"F:\1【机械零件】\Step官方标准数据\NIST-PMI-STEP-Files\nist_ftc_07_asme1_ap242-e2-sfa.xlsx"
XLSX = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_XLSX

# (group, handle, name, title, type, extra)
ROWS = [
 ("MBD_A",13706,"Simple Datum.1","A","composite_group_shape_aspect+composite_shape_aspect+datum_feature+shape_aspect",{"datum":"A"}),
 ("MBD_A",13747,"Simple Datum.2","B","composite_shape_aspect+datum_feature+shape_aspect",{"datum":"B"}),
 ("MBD_A",13788,"Simple Datum.3","C","composite_shape_aspect+datum_feature+shape_aspect",{"datum":"C"}),
 ("MBD_A",13928,"Geometrical Tolerance.1","⌓ .06 A B C","geometric_tolerance+geometric_tolerance_with_datum_reference+surface_profile_tolerance",{}),
 ("MBD_A",13959,"Profile tolerance of any surface.1","⌓ .015","surface_profile_tolerance",{}),
 ("MBD_A",13999,"Perpendicularity.1","⟂ Ø.015Ⓟ.425 A","perpendicularity_tolerance",{}),
 ("MBD_A",14040,"Position.1","⌖ Ø.02Ⓟ.425 A B","geometric_tolerance+geometric_tolerance_with_datum_reference+position_tolerance",{}),
 ("MBD_A",14101,"Position.2","⌖ Ø.023Ⓟ.425 A B C","geometric_tolerance+geometric_tolerance_with_datum_reference+position_tolerance",{}),
 ("MBD_A",14233,"Position.3","⌖ Ø.023Ⓟ.15 D B C","geometric_tolerance+geometric_tolerance_with_datum_reference+position_tolerance",{}),
 ("MBD_A",15146,"Linear Size.1","Ø0.250 +.003 / -0","dimensional_size",{"size type":"diameter"}),
 ("MBD_A",15209,"Oriented Linear Dimension.1","0.30 ±.02","dimensional_location",{}),
 ("MBD_A",15254,"Linear Size.2","Ø0.250 +.003 / -0","dimensional_size",{"size type":"diameter"}),
 ("MBD_A",15315,"Oriented Linear Dimension.2","0.30 ±.02","dimensional_location",{}),
 ("MBD_A",15360,"Linear Size.3","2× Ø0.250 +.003 / -0","dimensional_size",{"size type":"diameter"}),
 ("MBD_A",15405,"Linear Size.4","4× Ø0.250 +.003 / -0.001","dimensional_size",{"size type":"diameter"}),
 ("MBD_A",15506,"Oriented Linear Dimension.3","0.30 ±.02","dimensional_location",{}),
 ("MBD_A",15578,"Oriented Linear Dimension.4","0.30 ±.02","dimensional_location",{}),
 ("MBD_A",16261,"Text.1","Text.1","property_definition",{"property":"semantic text"}),
 ("MBD_A",16288,"Text.2","F1.2","property_definition",{"property":"semantic text"}),
 ("MBD_A",16316,"Text.3","F1.1","property_definition",{"property":"semantic text"}),
 ("MBD_A",16345,"Text.4","Text.4","property_definition",{"property":"semantic text"}),
 ("MBD_B",13837,"Simple Datum.4","D","composite_group_shape_aspect+composite_shape_aspect+datum_feature+shape_aspect",{"datum":"D"}),
 ("MBD_B",14143,"Position surfacic profile.1","⌓ .04 A B C .01 A","geometric_tolerance+geometric_tolerance_with_datum_reference+surface_profile_tolerance",{}),
 ("MBD_B",14269,"Flatness.1","⏥ .01","flatness_tolerance+geometric_tolerance+geometric_tolerance_with_defined_area_unit+geometric_tolerance_with_defined_unit",{}),
 ("MBD_B",14302,"Parallelism.1",".03 A","parallelism_tolerance",{}),
 ("MBD_B",14335,"Position surfacic profile.2","⌓ .06 A B C","geometric_tolerance+geometric_tolerance_with_datum_reference+surface_profile_tolerance",{}),
 ("MBD_B",14405,"Position.4","⌖ Ø.05Ⓢ A B C","geometric_tolerance+geometric_tolerance_with_datum_reference+geometric_tolerance_with_modifiers+position_tolerance",{"modifiers":"(.STATISTICAL_TOLERANCE.)"}),
 ("MBD_B",14498,"Position.5","⌖ Ø.06 A B C Ø.015 A","geometric_tolerance+geometric_tolerance_with_datum_reference+position_tolerance",{}),
 ("MBD_B",14558,"Position.6","⌖ Ø.04Ⓜ A B C","geometric_tolerance+geometric_tolerance_with_datum_reference+geometric_tolerance_with_modifiers+position_tolerance",{"modifiers":"(.MAXIMUM_MATERIAL_REQUIREMENT.)"}),
 ("MBD_B",14593,"Position.7","⌖ Ø.045 A B C","geometric_tolerance+geometric_tolerance_with_datum_reference+position_tolerance",{}),
 ("MBD_B",14735,"Position.8","⌖ Ø.06Ⓜ E B C Ø.02Ⓜ E B C","geometric_tolerance+geometric_tolerance_with_datum_reference+geometric_tolerance_with_modifiers+position_tolerance",{"modifiers":"(.MAXIMUM_MATERIAL_REQUIREMENT.)"}),
 ("MBD_B",15623,"Linear Size.5","3× Ø0.875 ±.02","dimensional_size",{"size type":"diameter","dimensional note":"statistical"}),
 ("MBD_B",15669,"Linear Size.6","3× Ø0.875 ±.01","dimensional_size",{"size type":"diameter"}),
 ("MBD_B",15715,"Linear Size.7","4× Ø0.173 +.005 / -0.001","dimensional_size",{"size type":"diameter"}),
 ("MBD_B",15761,"Linear Size.8","2× Ø0.594 +.012 / -0.002","dimensional_size",{"size type":"diameter"}),
 ("MBD_B",15807,"Linear Size.9","4× Ø0.281 +.006 / -0.001","dimensional_size",{"size type":"diameter"}),
 ("MBD_B",15899,"Linear Size.11","4× Ø0.438 ±.01","dimensional_size",{"size type":"diameter"}),
 ("MBD_B",15974,"Oriented Linear Dimension.5","0.30 ±.02","dimensional_location",{}),
 ("MBD_B",16374,"Text.5","Text.5","property_definition",{"property":"semantic text"}),
 ("MBD_B",16403,"Text.6","D","property_definition",{"property":"semantic text"}),
 ("MBD_B",16431,"Text.7","飦","property_definition",{"property":"semantic text"}),
 ("MBD_B",16490,"Text.9","鈱´","property_definition",{"property":"semantic text"}),
 ("MBD_C",13887,"Simple Datum.5","E","composite_group_shape_aspect+composite_shape_aspect+datum_feature+shape_aspect",{"datum":"E"}),
 ("MBD_C",14636,"Position surfacic profile.3","⌓ .07 A B C .015 A","geometric_tolerance+geometric_tolerance_with_datum_reference+surface_profile_tolerance",{}),
 ("MBD_C",14795,"Position.9","⌖ Ø.04Ⓜ A B C","geometric_tolerance+geometric_tolerance_with_datum_reference+geometric_tolerance_with_modifiers+position_tolerance",{"modifiers":"(.MAXIMUM_MATERIAL_REQUIREMENT.)"}),
 ("MBD_C",14830,"Position.10","⌖ Ø.045 A B C","geometric_tolerance+geometric_tolerance_with_datum_reference+position_tolerance",{}),
 ("MBD_C",16020,"Linear Size.12","2× Ø0.594 +.012 / -0.002","dimensional_size",{"size type":"diameter"}),
 ("MBD_C",16461,"Text.8","Text.8","property_definition",{"property":"semantic text"}),
 ("MBD_D",14866,"Position surfacic profile.5","⌓ .006 A B C","geometric_tolerance+geometric_tolerance_with_datum_reference+surface_profile_tolerance",{}),
 ("MBD_D",15050,"Position surfacic profile.6","⌓ .03 A B C .005 A B C","geometric_tolerance+geometric_tolerance_with_datum_reference+surface_profile_tolerance",{}),
 ("MBD_D",15098,"Position surfacic profile.7","⌓ .07 A B C .015 A","geometric_tolerance+geometric_tolerance_with_datum_reference+surface_profile_tolerance",{}),
 ("MBD_D",16200,"Linear Size.13","2× 0.012-0.022","dimensional_size",{"size type":"radius"}),
 ("MBD_D",16520,"Text.10","Text.10","property_definition",{"property":"semantic text"}),
]


def build_md() -> str:
    out = ["# PMI 提取结果", "", "- 视图分组：4 个", "- 标注总数：53 条", ""]
    cur = None
    for g, h, nm, title, ty, extra in ROWS:
        if g != cur:
            cur = g
            n = sum(1 for r in ROWS if r[0] == g)
            out += [f"## {g}", "", f"- 分组 Handle：`16{hash(g) % 900 + 100}`", f"- 标注数量：{n} 条", ""]
        seq = sum(1 for r in ROWS[:ROWS.index((g, h, nm, title, ty, extra))] if r[0] == g) + 1
        d = {"handle": str(h), "name": nm, "original_type_name": "draughting_callout", "type": ty}
        d.update(extra)
        out += [f"### {seq}. {title}", "", "detailData:", json.dumps(d, ensure_ascii=False, indent=2), ""]
    return "\n".join(out)


# ============================================================
# 断言工具
# ============================================================
PASS, FAIL = [], []


def check(name, got, want):
    if got == want:
        PASS.append(name)
    else:
        FAIL.append(f"{name}：期望 {want!r}，实际 {got!r}")


# ============================================================
# 一、纯函数单测
# ============================================================
# 1. 乱码修复（UTF-8 被按 GBK 解码）
check("乱码 鈱€ -> ⌀", C.fix_mojibake("鈱€.015"), "⌀.015")
check("乱码 鈱´ -> ⌴", C.fix_mojibake("鈱´"), "⌴")
check("正常中文不误伤", C.fix_mojibake("位置度"), "位置度")
check("正常中文不误伤2", C.fix_mojibake("机械零件"), "机械零件")
check("纯 ASCII 不变", C.fix_mojibake("Simple Datum.1"), "Simple Datum.1")

# 2. 数值归一（保留负号、忽略正号、消浮点噪声）
check("数值 +.003", C.norm_number("+.003"), "0.003")
check("数值 -.003", C.norm_number("-.003"), "-0.003")
check("数值 -.000", C.norm_number("-.000"), "-0")
check("数值 .02", C.norm_number(".02"), "0.02")
check("数值 浮点噪声", C.norm_number("0.2999999999"), "0.3")

# 3. 符号归一 + 分隔符保留
check("符号归一", C.normalize("Ø0.250 +.003 / -0"), "D0.25 0.003 / -0")
check("保留竖线分隔", C.normalize("A|B|C"), "A | B | C")
check("统计公差符号", C.normalize("飦"), "ST")
check("平面度符号", C.normalize("⏥ .01"), "▱ 0.01")

# 4. 语义指纹：同一公差的不同呈现方式应等价
t1, t2 = C.sem_tokens("⌖ Ø.04Ⓜ A B C"), C.sem_tokens("⌖ | ⌀0.040 Ⓜ | A | B | C")
check("指纹 数值", t1["numbers"], t2["numbers"])
check("指纹 基准", t1["letters"], t2["letters"])
check("指纹 修饰符", t1["modifiers"], t2["modifiers"])

# 5. 数量前缀（`4X`）不参与数值集合，单独比较
t3, t4 = C.sem_tokens("4× Ø0.173 +.005 / -0.001"), C.sem_tokens("4X ⌀0.173 +.005 -.001")
check("指纹 数量前缀剔除", t3["numbers"], t4["numbers"])
check("指纹 数量值", (t3["count"], t4["count"]), (4, 4))

# 6. 缺陷检测：乱码识别（含字节不全、无法程序化还原的残缺码）
check("乱码识别 鈱€", C.mojibake_chars("鈱€.015"), ["鈱", "€"])
check("乱码识别 鈱´", C.mojibake_chars("鈱´"), ["鈱", "´"])
check("乱码识别 飦（残缺映射表）", C.mojibake_chars("飦"), ["飦"])
check("乱码识别 正常中文不误伤", C.mojibake_chars("位置度"), [])
check("乱码识别 纯 ASCII", C.mojibake_chars("Simple Datum.1"), [])
check("乱码还原 鈱€ -> ⌀", C.repaired_text("鈱€.015"), "⌀.015")
check("乱码还原 飦 -> Ⓢ", C.repaired_text("飦"), "Ⓢ")

# 7. 数量前缀：SFA 图形文本里前缀不在串首（`DIM | 4X ⌀.250`），且常用 `×`
check("数量前缀 串首", C._CNT_PREFIX.search("4X ⌀0.173").group(1), "4")
check("数量前缀 竖线后", C._CNT_PREFIX.search("DIM | 4X ⌀.250").group(1), "4")
check("数量前缀 乘号写法", C._CNT_PREFIX.search("2× Ø0.438 ±.01").group(1), "2")
check("数量前缀 无", C._CNT_PREFIX.search("⌓ .06 A B C"), None)

# 8. 数值精度等价：SFA 图形通道对同一注释输出两行（全精度 + 展示精度）
check("精度等价 .4375/.438", C._num_equivalent("0.4375", "0.438"), True)
check("精度等价 .25/.250", C._num_equivalent("0.25", "0.250"), True)
check("精度等价 真差异", C._num_equivalent("0.015", "0.02"), False)
check("公差数值过滤 整数", C._tolerance_like("1"), False)
check("公差数值过滤 零", C._tolerance_like("0"), False)
check("公差数值过滤 真公差", C._tolerance_like("0.015"), True)

# ============================================================
# 二、markdown 解析
# ============================================================
from collections import Counter

md = build_md()
items = C.parse_dev_markdown(md)
check("解析条目数", len(items), 53)
cnt = Counter(i.kind for i in items)
check("几何公差条数", cnt[C.KIND_GT], 21)
check("尺寸公差条数", cnt[C.KIND_DIM], 17)
check("基准条数", cnt[C.KIND_DATUM], 5)
check("注释条数", cnt[C.KIND_NOTE], 10)
check("handle 解析", items[0].handle, 13706)
check("name 解析", items[0].name, "Simple Datum.1")
check("末条分组解析", items[-1].group, "MBD_D")

# ============================================================
# 三、集成测试（需真实 SFA 报告）
# ============================================================
if not os.path.exists(XLSX):
    print(f"[跳过集成测试] 未找到 SFA 报告：{XLSX}")
    print("   可执行： python test_pmi_core.py <你的SFA报告.xlsx>")
else:
    t = C.load_sfa(XLSX)
    # 工作表数只做「像不像一份完整报告」的粗筛，不锁死具体值：SFA 会把文件里出现的
    # 每一种实体类型都单独列一张表，同一测试件的不同导出差几十张是正常的
    # （实测 ftc_07 113~145、ftc_08 80~115）。锁死数值只会让换一份导出就红。
    check("SFA 工作表数（粗筛）", len(t.sheets) > 50, True)
    check("SFA 语义项数", len(t.semantic), 49)
    check("SFA draughting_callout 行数", len(t.dc), 53)
    check("SFA 图形注释行数", len(t.ta), 53)
    check("SFA datum 行数", len(t.datum), 5)
    check("无解析警告", t.warnings, [])

    # 关联键：md.handle 必须等于 SFA draughting_callout.ID，且 name 一致
    check("全部 handle 命中 SFA", [i.handle for i in items if i.handle not in t.dc], [])
    check("handle/name 与 SFA 一致",
          [i.name for i in items if i.handle in t.dc and t.dc[i.handle] != i.name], [])

    rows = C.match_items(t, items)
    m = C.compute_metrics(rows)

    # 结果表与汇总表必须带缺陷字段
    check("结果表含缺陷列", [c for c in rows[0].as_dict() if "缺陷" in c], ["提取缺陷", "缺陷详情"])
    s = C.summarize(rows, t, items)
    check("汇总含缺陷字段",
          [k for k in s if k.startswith("缺陷_")],
          ["缺陷_条目数", "缺陷_总数", "缺陷_占比", "缺陷_分布"])

    # 三条关联路径都必须出现，且无漏网
    paths = Counter(r.path for r in rows)
    check("1 跳直连（几何公差）", paths["gt-1hop"], 27)
    check("2 跳（尺寸 dcr）", paths["dim-2hop"], 17)
    check("2 跳（基准 datum）", paths["datum-2hop"], 5)
    check("注释条目", paths["note"], 10)
    check("基准体系组合项", paths["datum_system"], 5)
    check("无未关联条目", [r.key for r in rows if r.path == "none"], [])

    # 三层指标
    check("语义 应提取", m.sem_expected, 44)
    check("语义 缺失", m.sem_miss, 0)
    check("语义 多余", m.sem_extra, 0)
    check("语义 疑似", m.sem_suspect, 0)
    check("语义 召回率", round(m.recall, 2), 100.00)
    check("语义 精确率", round(m.precision, 2), 100.00)
    check("基准 应提取", m.datum_expected, 5)
    check("基准 命中", m.datum_hit, 5)
    check("注释 总数", m.note_total, 10)
    check("注释 图形通道有文本", m.note_graphic_only, 3)
    check("注释 内部独占", m.note_exclusive, 7)

    # 人工校验不污染自动指标
    m2 = C.compute_metrics(rows, {r.key: "待定" for r in rows})
    check("初始/待定指标一致", (m.sem_hit, round(m.recall, 2)), (m2.sem_hit, round(m2.recall, 2)))

    # 已知差异应被正确识别为「文本差异」而非「疑似」
    diff_kinds = {r.key: r.status for r in rows if r.status == C.ST_HIT_DIFF}
    check("复合公差合并判定",
          diff_kinds.get("MBD_B#2@14121"), C.ST_HIT_DIFF)
    check("数量前缀差异判定",
          diff_kinds.get("MBD_D#4@16189"), C.ST_HIT_DIFF)

    # ------------------------------------------------------------
    # 四、开发侧提取缺陷层
    # 缺陷=开发侧数据本身错（乱码/丢符号/丢前缀），必须报出且不得影响判定
    # ------------------------------------------------------------
    # 与指标层用同一套聚合（同一开发标注的多个子行合并去重）：
    # 逐子行取第一条会漏掉「第一子行无码、第二子行才有」的缺陷，
    # `MRG` 正是这类 —— 它只在数值多出的那些子行上判得出来。
    defects = {k: (g["codes"], g["detail"]) for k, g in C.defect_groups(rows).items()}

    check("缺陷 条目数与指标一致", m.defect_rows, len(defects))
    check("缺陷 ENC 编码损坏条目",
          sorted(k for k, (c, _) in defects.items() if C.DF_ENC in c), ["MBD_B#20", "MBD_B#21"])
    check("缺陷 ENC 明细给出还原值", "Ⓢ" in defects["MBD_B#20"][1], True)
    check("缺陷 CNT 数量前缀不符条目",
          sorted(k for k, (c, _) in defects.items() if C.DF_CNT in c),
          ["MBD_A#16", "MBD_A#17", "MBD_B#17", "MBD_D#4"])
    check("缺陷 SYM 符号丢失 (Parallelism.1 缺 ∥)",
          C.DF_SYM in defects.get("MBD_B#4", ([], ""))[0], True)
    check("缺陷 SYM 符号丢失 (Flatness.1 缺 ⌀ 定义区域)",
          C.DF_SYM in defects.get("MBD_B#3", ([], ""))[0], True)
    check("缺陷 NUM 无误报", C.DF_NUM in m.defect_by_code, False)
    check("缺陷 图形通道双行精度差不计入缺陷", "MBD_A#15" in defects, False)
    check("缺陷 无 MAP / DUP 误报",
          [k for k, (c, _) in defects.items() if C.DF_MAP in c or C.DF_DUP in c], [])
    # MRG 是唯一的「对照真值类」缺陷码：夹具里这 6 条都是开发侧把 SFA 拆开的
    # 复合公差（轮廓度 ⌓ / 位置度 ⌖）并成一条，判据是「映射到多个语义实体
    # 且开发侧数值是超集」——两条缺一不可，缺前者会把「单纯多提一个数值」也算进来。
    check("缺陷 MRG 复合公差未拆分条目",
          sorted(k for k, (c, _) in defects.items() if C.DF_MRG in c),
          ["MBD_B#10", "MBD_B#2", "MBD_B#7", "MBD_C#2", "MBD_D#2", "MBD_D#3"])
    check("缺陷 MRG 明细点名 SFA 拆分条数",
          "SFA 拆成 2 条" in defects["MBD_B#2"][1], True)
    check("缺陷 MRG 条目均映射多个语义实体（判据前提）",
          sorted(k for k in defects if C.DF_MRG in defects[k][0]
                 and len([r for r in rows if r.key.split("@")[0] == k]) < 2), [])
    check("缺陷 分布", m.defect_by_code,
          {C.DF_CNT: 4, C.DF_MRG: 6, C.DF_SYM: 3, C.DF_ENC: 2})
    check("缺陷 总数", m.defect_total, 15)

    # 缺陷清单（展示 / 导出用）：必须带 SFA 侧原值 —— 只看开发侧文本没法判断
    # 「到底哪对不上」，排查时还得回查报告。
    recs = C.defect_records(rows)
    check("缺陷清单 条数与指标一致", len(recs), m.defect_rows)
    check("缺陷清单 字段齐全",
          set(recs[0]) >= {"分组", "开发名称", "开发标注", "Handle", "缺陷",
                           "SFA条数", "SFA语义ID", "SFA语义文本", "SFA图形文本",
                           "关联路径", "详情", "定位"}, True)
    check("缺陷清单 至少给出一条 SFA 原值",
          any(r["SFA语义文本"] for r in recs), True)
    _mrg_rec = next(r for r in recs if C.DF_MRG in r["缺陷"])
    check("缺陷清单 复合公差给出 SFA 多值与条数",
          (_mrg_rec["SFA条数"], len(_mrg_rec["SFA语义文本"].split("／"))), (2, 2))
    check("缺陷清单 多值去重（同一 FCF 不重复列）",
          len(set(_mrg_rec["SFA语义文本"].split("／"))), 2)
    check("缺陷清单 带关联路径与 handle",
          bool(_mrg_rec["关联路径"]) and _mrg_rec["Handle"] != "", True)
    # 缺陷与分析指标解耦：有缺陷的条目只要内容对得上，仍计入命中
    check("缺陷不拉低指标（缺陷条目仍全命中）",
          [r.key for r in rows if r.defects and r.layer == C.LAYER_SEM
           and r.status not in (C.ST_HIT, C.ST_HIT_DIFF)], [])
    check("缺陷不进 precision 分母", m.sem_extracted, 44)

    print("\n--- 开发侧提取缺陷（本工具要产出的核心结论）---")
    for k, (codes, detail) in sorted(defects.items()):
        print(f"  {k:10s} {' '.join(codes):34s} {detail}")
    print("  缺陷分布:", m.defect_by_code)
    print("  缺陷占比: %.2f%%（%d 条开发标注含缺陷 / 共 %d 条参与比对）"
          % (m.defect_rate, m.defect_rows,
             m.sem_extracted + m.datum_expected + m.note_total))

# ============================================================
# 三、解析自检（ParseDiag）—— 覆盖「一条都没对上」的各类畸形输入
# ============================================================
def _md(detail_block: str, heading: str = "### 1. 标题A") -> str:
    return f"## MBD_A\n\n- 标注数量：1 条\n\n{heading}\n\n{detail_block}\n"


_GOOD = '{\n  "handle": "13706",\n  "name": "Simple Datum.1",\n  "type": "datum_feature"\n}'

# 正常
it, dg = C.parse_dev_markdown_ex(_md("detailData:\n" + _GOOD))
check("自检 正常 条目数", dg.items, 1)
check("自检 正常 带handle", dg.with_handle, 1)
check("自检 正常 带name", dg.with_name, 1)
check("自检 正常 healthy", dg.healthy, True)
check("自检 正常 无问题", dg.problems(), [])

# 标签与 { 同行（旧实现会静默失败）
it, dg = C.parse_dev_markdown_ex(_md("detailData: " + _GOOD))
check("自检 detailData 同行 仍能解析", dg.with_handle, 1)
check("自检 detailData 同行 healthy", dg.healthy, True)

# 带 ```json 围栏
it, dg = C.parse_dev_markdown_ex(_md("detailData:\n```json\n" + _GOOD + "\n```"))
check("自检 围栏 仍能解析", dg.with_handle, 1)

# 全角冒号 + 引号标签
it, dg = C.parse_dev_markdown_ex(_md('"detailData"：\n' + _GOOD))
check("自检 全角冒号+引号 仍能解析", dg.with_handle, 1)

# 缺 detailData 块 -> 必然全判「多余」
it, dg = C.parse_dev_markdown_ex(_md(""))
check("自检 缺块 条目数", dg.items, 1)
check("自检 缺块 记为未解析", len(dg.detail_missing), 1)
check("自检 缺块 未取到 handle", dg.with_handle, 0)
check("自检 缺块 healthy", dg.healthy, False)
check("自检 缺块 给出提示", any("detailData" in p for p in dg.problems()), True)

# broken JSON
it, dg = C.parse_dev_markdown_ex(_md('detailData:\n{\n  "handle": "13706",\n')) 
check("自检 JSON 截断 记为失败", len(dg.json_failed), 1)
check("自检 JSON 截断 未取到 handle", dg.with_handle, 0)

# 字段名大小写
it, dg = C.parse_dev_markdown_ex(
    _md('detailData:\n{\n  "Handle": "13706",\n  "Name": "Simple Datum.1"\n}'))
check("自检 大写字段名 容错", (dg.with_handle, dg.with_name), (1, 1))

# JSON 在但缺 handle/name
it, dg = C.parse_dev_markdown_ex(_md('detailData:\n{\n  "type": "datum_feature"\n}'))
check("自检 缺关键字段 记为 key_missing", len(dg.key_missing), 1)

# 标题没序号
it, dg = C.parse_dev_markdown_ex(_md("detailData:\n" + _GOOD, heading="### 标题A"))
check("自检 标题无序号 条目数", dg.items, 0)
check("自检 标题无序号 记录未匹配行", len(dg.headings_unmatched), 1)
check("自检 标题无序号 给出提示",
      any("没有解析出任何标注条目" in p for p in dg.problems()), True)

# 四级标题（解析不出条目）
it, dg = C.parse_dev_markdown_ex(_md("detailData:\n" + _GOOD, heading="#### 1. 标题A"))
check("自检 四级标题 条目数", dg.items, 0)
check("自检 四级标题 记录未匹配行", len(dg.headings_unmatched), 1)

# 兼容入口仍返回纯条目
check("parse_dev_markdown 仍返回条目", len(C.parse_dev_markdown(_md("detailData:\n" + _GOOD))), 1)


# ============================================================
# 四、SFA 列数容错 / datum_target / 分段取文本 / 修饰符字形
#     覆盖「一条都没比对上」的根因：
#       - tessellated_annotation_occurrence 缺 col14 时整表被跳过 -> 关联全断
#       - datum_target 未走 1 跳引用 -> 基准目标全判多余
#       - 语义行把尺寸行与 FCF 行揉在一起 -> 内容比对大面积误报
# ============================================================
def _build_sfa(tmpdir: str) -> str:
    """构造一份最小 SFA 报告：ta 表刻意只给 12 列（缺 Equivalent Unicode String）。"""
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Header")
    for r in ([f"Header row {i}", None, None] for i in range(3)):
        ws.append(r)

    # 语义表：注意 2002 号把「尺寸行 + FCF 行」揉在一行
    ws = wb.create_sheet("Semantic PMI Summary")
    ws.append(["Semantic PMI Summary (5)", None, None, None])
    ws.append(["ID", "Entity", "Semantic PMI", "Similar"])
    for rid, ent, txt in [
        (1001, "datum_system", "A | B"),
        (2001, "position_tolerance", "⌖ | ⌀0.8 | A | B"),
        (2002, "cylindricity_tolerance", "2X ⌀3.50 ± 0.2\n⌭ | 0.1"),
        (2003, "datum_target", "K1 (area)"),
    ]:
        ws.append([str(rid), ent, txt, None])

    ws = wb.create_sheet("draughting_callout")
    ws.append(["draughting_callout  (4)", None])
    for rid, nm in [(1001, "Simple Datum.1"), (1101, "Position.1"),
                    (1201, "Cylindricity.1"), (1301, "Datum Target.1")]:
        ws.append([str(rid), nm])

    # ta 表只写 12 列，第 12 列（Associated Semantic PMI）在索引 10
    ws = wb.create_sheet("tessellated_annotation_occurren")
    ws.append(["tessellated_annotation_occurrence (4)"] + [None] * 11)
    ws.append(["ID", "name", "styles", "item", "name", "children",
               "presentation style", "color", "plane", "Associated Geometry",
               "Associated Semantic PMI", "Saved Views"])
    for rid, nm, ref in [(1100, "Position.1", "position_tolerance 2001"),
                         (1200, "Cylindricity.1", "cylindricity_tolerance 2002"),
                         (1000, "Simple Datum.1", "datum_feature 1002"),
                         (1300, "Datum Target.1", "datum_target 2003")]:
        ws.append([str(rid), nm] + [""] * 8 + [ref, ""])

    ws = wb.create_sheet("datum")
    ws.append(["datum  (1)", None, None, None, None, None])
    # 只放 A：`Simple Datum.1` 的 ta 引用是 1002，靠 ID-1 两跳命中 1001
    for rid, ident in [(1001, "A")]:
        ws.append([str(rid), "datum", "x", "x", "x", ident])

    path = os.path.join(tmpdir, "mini_sfa.xlsx")
    wb.save(path)
    return path


_MD_MINI = """# PMI 提取结果

## MBD_X

- 标注数量：4 条

### 1. A

detailData:
{
  "handle": "1001",
  "datum": "A",
  "name": "Simple Datum.1",
  "type": "datum_feature"
}

### 2. ⌖ Ø.8 A B

detailData:
{
  "handle": "1101",
  "name": "Position.1",
  "type": "position_tolerance"
}

### 3. .1

detailData:
{
  "handle": "1201",
  "name": "Cylindricity.1",
  "type": "cylindricity_tolerance"
}

### 4. K1

detailData:
{
  "handle": "1301",
  "datum": "K",
  "target id": "1",
  "name": "Datum Target.1",
  "type": "datum_target"
}
"""

import tempfile  # noqa: E402

with tempfile.TemporaryDirectory() as _td:
    _truth = C.load_sfa(_build_sfa(_td))

check("列数容错 ta 表 12 列仍装载", len(_truth.ta), 4)
check("列数容错 ta 列数已记录", _truth.ta_cols, 12)
check("列数容错 缺图形文本通道时给出告警", bool(_truth.warnings), True)
check("列数容错 告警指向 Unicode 列",
      any("Equivalent Unicode String" in w for w in _truth.warnings), True)

_items_mini, _ = C.parse_dev_markdown_ex(_MD_MINI)
_rows_mini = C.match_items(_truth, _items_mini)
_m_mini = C.compute_metrics(_rows_mini)
check("列数容错 全部关联成功（无缺失）", _m_mini.sem_miss + _m_mini.sem_extra, 0)
check("列数容错 召回 100", _m_mini.recall, 100.0)
check("列数容错 基准 100", _m_mini.datum_coverage, 100.0)
check("datum_target 走 1 跳引用",
      C._lookup_semantic_for_dev(_items_mini[3], _truth)[1], "datum_target-1hop")
check("datum_target 判定为命中",
      [r.status for r in _rows_mini if r.key.startswith("MBD_X#4")], [C.ST_HIT])
check("普通基准走 ID-1 两跳",
      C._lookup_semantic_for_dev(_items_mini[0], _truth)[1], "datum-2hop")

# GT 分段：语义行含「尺寸行 + FCF 行」时只取 FCF 行，且保留原文字形
check("分段 GT 只取 FCF 行",
      C.sfa_text_for("2X ⌀3.50 ± 0.2\n⌭ | 0.1", C.KIND_GT), "⌭ | 0.1")
check("分段 GT 保留原文字形",
      C.sfa_text_for("⌖ | ⌀0.8 | A | B", C.KIND_GT), "⌖ | ⌀0.8 | A | B")
check("分段 无 FCF 行时退回首行",
      C.sfa_text_for("⌀3.50 G6", C.KIND_GT), "⌀3.50 G6")
check("分段 DIM 合并折行",
      C.sfa_text_for("⌀0.250 +.003\n-.000", C.KIND_DIM), "⌀0.250 +.003 -.000")

# 数量前缀不得污染字母集（`2X` 的 X 曾被当成基准字母）
_tk_cnt = C.sem_tokens("2X ⌀5.50 ± 0.2")
check("数量前缀 已识别", _tk_cnt["count"], 2)
check("数量前缀 不泄漏进基准字母", "X" in _tk_cnt["letters"], False)

# 修饰符优先按字形识别（开发 `EⓂ-FⓂ-GⓂ` vs SFA `E Ⓜ-F Ⓜ-G Ⓜ`）
_tk_dev = C.sem_tokens("⌖ Ø.8 D EⓂ-FⓂ-GⓂ")
_tk_sfa = C.sem_tokens("⌖ | ⌀0.8 | D | E Ⓜ-F Ⓜ-G Ⓜ")
check("修饰符字形 开发侧识别 M", _tk_dev["modifiers"], ["M"])
check("修饰符字形 SFA 侧识别 M", _tk_sfa["modifiers"], ["M"])
check("修饰符字形 基准字母含 F", _tk_dev["letters"], ["D", "E", "F", "G"])
check("修饰符字形 两侧判定一致",
      C.compare_tokens(_tk_dev, _tk_sfa)[0], C.ST_HIT)

# 基准字母与修饰符同形（`L` 既是基准又是 LMC）需对称消歧
check("同形字母 对称剔除后判命中",
      C.compare_tokens(C.sem_tokens("⌖ Ø0 H K L"),
                       C.sem_tokens("⌖ | ⌀0 Ⓛ | H | K | L"))[0], C.ST_HIT)

# 符号集：直线度用 U+2212、圆柱度 ⌭ 要检出；`⭩◎` 属渲染杂符不得误报
check("符号 直线度符已收录", C.symbols_of("− | 0.2 / 15"), {"−"})
check("符号 圆柱度符已收录", C.symbols_of("⌭ | 0.1"), {"⌭"})
check("符号 杂符 ⭩◎ 不计入", C.symbols_of("⭩◎ | ⌓ | 1.5"), {"⌓"})

# SFA 把统计公差渲染成字体私有区字形 U+F055（同一份报告里另一种写法是 `<ST>`），
# 认不出就等于丢掉修饰符
check("符号 私有区 ST 字形归一化为 ST", C.normalize("\uf055"), "ST")
check("符号 指向符号 ↧ 被剔除", "↧" in C.normalize("DIM | ↧.30±.02"), False)
check("符号 注释杂符 ⌴ 被剔除", C.normalize("⌴"), "")
# ▽ / ⎹ 出现在带基准框架里、与 [A] 同行，属 SFA 版面元素；
# 混进公差符号集会让「缺符号」检查误报一片
check("符号 排版字形不混入公差符号集",
      C.symbols_of("▽ ⎹ ⌓ | 1.2 | A"), {"⌓"})
check("符号 排版字形已登记", {"▽", "⎹"} <= C.SFA_LAYOUT_GLYPHS, True)


# ============================================================
# 五、列名驱动装载 / 加载期交叉校验 / 断链候选
#     覆盖本次事故的同源根因：装载器只认列号，列数或列序一变就静默失配
#     （曾整张 tessellated 表被跳过 -> 结果全判「多余 / 缺失」）。
#     表头列名其实就在表里（带换行与 (Sec. x) 后缀），以前从来没读。
# ============================================================
def _build_sfa_shuffled(tmpdir: str) -> str:
    """列序全部打乱 + 表头带换行与 (Sec. x) 后缀的报告。

    若装载器仍按列号取，ref 会落到空列、关联全断；按列名定位则应完全正常。
    """
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Semantic PMI Summary")
    ws.append(["Semantic PMI Summary (3)", None, None, None])
    # 表头带换行与章节号；Expected PMI 被挪到第 4 列
    ws.append(["Entity\n(Sec. 2.1)", "ID", "Similar PMI / Exception", "Semantic PMI"])
    for rid, ent, txt in [(2001, "position_tolerance", "⌖ | ⌀0.8 | A | B"),
                          (3001, "flatness_tolerance", "▱ | 0.2"),
                          (4001, "dimensional_size", "⌀6.00 ± 0.15")]:
        ws.append([ent, str(rid), None, txt])

    ws = wb.create_sheet("draughting_callout")
    ws.append(["draughting_callout  (3)", None])
    ws.append(["name", "ID"])                       # 列序反转
    for rid, nm in [(1101, "Position.1"), (1201, "Flatness.1"), (1301, "Linear Size.1")]:
        ws.append([nm, str(rid)])

    ws = wb.create_sheet("tessellated_annotation_occurren")
    ws.append(["tessellated_annotation_occurrence  (3)"] + [None] * 4)
    ws.append(["列1", "列2", "列3", "列4", "列5"])    # Excel 占位行，须跳过
    ws.append(["ID", "name", "Associated Semantic PMI (Sec. 7.3)", "styles",
               "Equivalent Unicode String(s)\n(Sec. 10.1.3.3)"])
    for rid, nm, ref, uni in [
        (1100, "Position.1", "position_tolerance 2001", "⌖ | ⌀0.8 | A | B"),
        (1200, "Flatness.1", "flatness_tolerance 3001", "▱ | 0.2"),
        (1300, "Linear Size.1", "dimensional_size 5001", "⌀6.00 ± 0.15"),
    ]:
        ws.append([str(rid), nm, ref, "", uni])

    ws = wb.create_sheet("dimensional_characteristic_repr")
    ws.append(["dimensional_characteristic_representation  (1)"] + [None] * 2)
    ws.append(["ID", "dimension (Sec. 5.1.1)", "Associated Geometry"])
    ws.append(["4001", "dimensional_size 5001", "x"])

    path = os.path.join(tmpdir, "shuffled_sfa.xlsx")
    wb.save(path)
    return path


def _build_sfa_noheader(tmpdir: str) -> str:
    """ta 表没有可用表头（只有 Excel 占位行），且 ref 列不在默认位置。

    期望：退回默认列号、标记为不可信、交叉校验报出来 —— 而不是静默失配。
    """
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Semantic PMI Summary")
    ws.append(["Semantic PMI Summary (1)", None, None, None])
    ws.append(["ID", "Entity", "Semantic PMI", "Similar"])
    ws.append(["2001", "position_tolerance", "⌖ | ⌀0.8 | A | B", None])

    ws = wb.create_sheet("draughting_callout")
    ws.append(["draughting_callout  (1)", None])
    ws.append(["ID", "name"])
    ws.append(["1101", "Position.1"])

    ws = wb.create_sheet("tessellated_annotation_occurren")
    ws.append(["tessellated_annotation_occurrence  (1)"] + [None] * 4)
    ws.append(["列1", "列2", "列3", "列4", "列5"])     # 唯一「表头」就是占位行
    ws.append(["1100", "Position.1", "", "position_tolerance 2001", ""])

    path = os.path.join(tmpdir, "noheader_sfa.xlsx")
    wb.save(path)
    return path


def _build_sfa_nonumeric(tmpdir: str) -> str:
    """模拟本次事故形态：ta 表 ID 列读不出数字，整表装载为 0。

    期望：交叉校验直接报「无可识别的数据行」，而不是等到结果表全红才发现。
    """
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Semantic PMI Summary")
    ws.append(["Semantic PMI Summary (1)", None, None, None])
    ws.append(["ID", "Entity", "Semantic PMI", "Similar"])
    ws.append(["2001", "position_tolerance", "⌖ | ⌀0.8 | A | B", None])

    ws = wb.create_sheet("draughting_callout")
    ws.append(["draughting_callout  (1)", None])
    ws.append(["ID", "name"])
    ws.append(["1101", "Position.1"])

    ws = wb.create_sheet("tessellated_annotation_occurren")
    ws.append(["tessellated_annotation_occurrence  (1)"] + [None] * 11)
    ws.append(["ID", "name", "styles", "item", "name", "children", "presentation style",
               "color", "plane", "Associated Geometry", "Associated Semantic PMI",
               "Saved Views"])
    # ID 带 `#` 前缀 -> isdigit() 为假 -> 整表装载 0
    ws.append(["#1100", "Position.1", "", "", "", "", "", "", "", "",
               "position_tolerance 2001", ""])

    path = os.path.join(tmpdir, "nonumeric_sfa.xlsx")
    wb.save(path)
    return path


_MD_SHUFFLED = """# PMI 提取结果

## MBD_Y

- 标注数量：3 条

### 1. ⌖ Ø.8 A B

detailData:
{
  "handle": "1101",
  "name": "Position.1",
  "type": "position_tolerance"
}

### 2. ⏥ .2

detailData:
{
  "handle": "1201",
  "name": "Flatness.1",
  "type": "flatness_tolerance"
}

### 3. Ø6.00 ±.15

detailData:
{
  "handle": "1301",
  "name": "Linear Size.1",
  "type": "dimensional_size"
}
"""

with tempfile.TemporaryDirectory() as _td_sh:
    _sh = C.load_sfa(_build_sfa_shuffled(_td_sh))
_sh_items, _ = C.parse_dev_markdown_ex(_MD_SHUFFLED)
_sh_rows = C.match_items(_sh, _sh_items)
_sh_m = C.compute_metrics(_sh_rows)

check("列序打乱 语义表仍装载 3 项", len(_sh.semantic), 3)
check("列序打乱 dc 表 ID/name 反转仍装载", len(_sh.dc), 3)
check("列序打乱 ta 表 ref 列挪位仍装载", len(_sh.ta), 3)
check("列序打乱 dcr 表仍装载", len(_sh.dcr_by_dim), 1)
check("列序打乱 语义表定位到 Semantic PMI 列",
      _sh.recipe("semantic").resolved["semantic pmi"], 3)
check("列序打乱 dc 定位到 name=0 / id=1",
      (_sh.recipe("dc").resolved["name"], _sh.recipe("dc").resolved["id"]), (0, 1))
check("列序打乱 ta 定位到 ref 列 2",
      _sh.recipe("ta").resolved["associated semantic pmi"], 2)
check("列序打乱 ta 识别出 Unicode 列",
      _sh.recipe("ta").resolved["equivalent unicode string"], 4)
check("列序打乱 每张表都按表头识别（无 fallback）",
      sorted({r.source for r in _sh.recipes}), ["header"])
check("列序打乱 交叉校验全通过", [b.line() for b in _sh.broken_checks()], [])
check("列序打乱 召回 100", _sh_m.recall, 100.0)
check("列序打乱 精确 100", _sh_m.precision, 100.0)
check("列序打乱 无断链条目", C.link_stats(_sh_rows).get("none", 0), 0)

with tempfile.TemporaryDirectory() as _td_nh:
    _nh = C.load_sfa(_build_sfa_noheader(_td_nh))
check("无表头 ta 定位不可信", _nh.recipe("ta").trustworthy, False)
check("无表头 ta 仍装载（不静默跳过整表）", len(_nh.ta), 1)
check("无表头 ta ref 取空（关联会失效）", _nh.ta["Position.1"]["sem_refs"], [])
check("无表头 交叉校验报出列名未识别",
      any("列名识别" in b.name for b in _nh.broken_checks()), True)
check("无表头 语义表不受影响、仍可靠",
      _nh.recipe("semantic").trustworthy, True)

with tempfile.TemporaryDirectory() as _td_ne:
    _ne = C.load_sfa(_build_sfa_nonumeric(_td_ne))
check("ID 非数字 ta 装载 0 条", len(_ne.ta), 0)
check("ID 非数字 交叉校验报出无可识别数据行",
      any("无可识别的数据行" in b.value for b in _ne.broken_checks()), True)

# 断链候选：只提示、不参与判定
_MD_TYPO = _MD_SHUFFLED.replace('"name": "Position.1"', '"name": "Position.11"')
_tp_items, _ = C.parse_dev_markdown_ex(_MD_TYPO)
_tp_rows = C.match_items(_sh, _tp_items)
_tp_m = C.compute_metrics(_tp_rows)
_sus = C.suspect_links(_tp_items, _sh, _tp_rows)

check("断链候选 关联确已断裂", C.link_stats(_tp_rows).get("none", 0), 1)
check("断链候选 给出候选", len(_sus), 1)
check("断链候选 指向正确条目",
      [c["ta_name"] for c in _sus[0].candidates] if _sus else [], ["Position.1"])
check("断链候选 判定未被改写（条目仍判多余）", _tp_m.sem_extra >= 1, True)
check("断链候选 判定未被改写（对端仍判缺失）", _tp_m.sem_miss >= 1, True)
check("断链候选 正常数据下不产生噪音",
      C.suspect_links(_sh_items, _sh, _sh_rows), [])


# ============================================================
# CTC 系列特征：短 ID + datum_feature 独立通道（NIST ctc_01 踩出的三个坑）
#   a) 实体 ID 只有 2~3 位（FTC 是 4~7 位），引用正则写死位数会让 dcr 整表装载 0 条
#   b) datum 与 datum_feature 是两套 ID，不保证相邻，不能靠 `ID-1` 推
#   c) 尺寸的 dcr 映射不能做 `ID+1` 兜底，否则会串到相邻尺寸
# ============================================================
def _build_sfa_ctc_like(tmpdir: str) -> str:
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Header")
    for i in range(3):
        ws.append([f"Header row {i}", None])

    ws = wb.create_sheet("Semantic PMI Summary")
    ws.append(["Semantic PMI Summary (2)", None, None, None])
    ws.append(["ID", "Entity", "Semantic PMI", "Similar PMI"])
    for rid, ent, txt in [
        (51, "datum_system", "A"),
        (113, "dimensional_characteristic_representation", "⌀25 ± 0.15"),
        (118, "dimensional_characteristic_representation", "60° ± 0.5°"),
    ]:
        ws.append([str(rid), ent, txt, None])

    ws = wb.create_sheet("draughting_callout")
    ws.append(["draughting_callout  (4)", None])
    ws.append(["ID", "name"])
    for rid, nm in [(607, "Linear Size.1"), (616, "Linear Size.5"),
                    (618, "Angular Size.1"), (615, "Simple Datum.1")]:
        ws.append([str(rid), nm])

    ws = wb.create_sheet("tessellated_annotation_occurren")
    ws.append(["tessellated_annotation_occurrence (4)"] + [None] * 11)
    ws.append(["ID", "name", "styles", "item", "name", "children",
               "presentation style", "color", "plane", "Associated Geometry",
               "Associated Semantic PMI", "Saved Views"])
    for rid, nm, ref in [(600, "Linear Size.1", "dimensional_size 120"),
                         (601, "Linear Size.5", "dimensional_size 121"),
                         (602, "Angular Size.1", "angular_location 33"),
                         (603, "Simple Datum.1", "datum_feature 34")]:
        ws.append([str(rid), nm] + [""] * 8 + [ref, ""])

    # datum_feature 34 -> A，而 datum 是 37 -> A：差 3，不是 1
    ws = wb.create_sheet("datum_feature")
    ws.append(["datum_feature  (1)"] + [None] * 5)
    ws.append(["ID", "name", "description", "of_shape", "product_definitional", "Datum"])
    ws.append(["34", "Simple Datum.1", "", "", "True", "A"])

    ws = wb.create_sheet("datum")
    ws.append(["datum  (1)"] + [None] * 5)
    ws.append(["ID", "name", "description", "of_shape", "product_definitional", "identification"])
    ws.append(["37", "", "", "", "False", "A"])

    # dcr：121 -> 113（Linear Size.5）、33 -> 118（Angular Size.1）；**120 故意没有**
    ws = wb.create_sheet("dimensional_characteristic_repr")
    ws.append(["dimensional_characteristic_representation (2)"] + [None] * 3)
    ws.append(["ID", "dimension", "representation", "Dimensional Tolerance"])
    ws.append(["113", "dimensional_size 121", "", "⌀25 ± 0.15"])
    ws.append(["118", "angular_location 33", "", "60° ± 0.5°"])

    # 无语义文本的实体表：120 / 121 存在，只是没有 dcr
    ws = wb.create_sheet("dimensional_size")
    ws.append(["dimensional_size  (2)"])
    ws.append(["ID", "applies_to", "name"])
    ws.append(["120", "composite_shape_aspect 219", "diameter"])
    ws.append(["121", "composite_shape_aspect 220", "diameter"])

    path = os.path.join(tmpdir, "ctc_like_sfa.xlsx")
    wb.save(path)
    return path


_MD_CTC = """# PMI 提取结果

## MBD_0

- 标注数量：4 条

### 1. Ø35 0 / -0.2

detailData:
{
  "handle": "607",
  "name": "Linear Size.1",
  "size type": "diameter",
  "type": "dimensional_size"
}

### 2. Ø25 ±.15

detailData:
{
  "handle": "616",
  "name": "Linear Size.5",
  "size type": "diameter",
  "type": "dimensional_size"
}

### 3. 60 ±.5

detailData:
{
  "handle": "618",
  "length type": "angle",
  "name": "Angular Size.1",
  "type": "angular_location"
}

### 4. A

detailData:
{
  "handle": "615",
  "datum": "A",
  "name": "Simple Datum.1",
  "type": "datum_feature"
}
"""

with tempfile.TemporaryDirectory() as _td_ctc:
    _ctc = C.load_sfa(_build_sfa_ctc_like(_td_ctc))
_ctc_items, _ = C.parse_dev_markdown_ex(_MD_CTC)
_ctc_rows = C.match_items(_ctc, _ctc_items)
_ctc_m = C.compute_metrics(_ctc_rows)
_ctc_paths = C.link_stats(_ctc_rows)

check("短 ID 引用可提取（3 位数不丢）", _ctc.ta["Linear Size.1"]["sem_refs"], [120])
check("短 ID dcr 表不再整表跳过", _ctc.recipe("dcr").n_loaded, 2)
check("短 ID 交叉校验全通过", [b.line() for b in _ctc.broken_checks()], [])

check("datum_feature 通道优先于 ID 邻接",
      C._lookup_semantic_for_dev(_ctc_items[3], _ctc)[1], "datum_feature-1hop")
check("datum_feature 基准字母取对", _ctc.datum_feature.get(34), "A")
check("基准覆盖率 100", _ctc_m.datum_coverage, 100.0)

check("尺寸不借 ID+1 兜底（120 无 dcr 即判无关联）",
      C._lookup_semantic_for_dev(_ctc_items[0], _ctc), ([], "none"))
check("无语义值的尺寸归注释层", _ctc_m.note_exclusive >= 1, True)
# 语义分母只算「两边都有内容可比」的条目：#2 尺寸 + #3 角度；#4 基准单列一层
check("无语义值的尺寸不计入语义分母", _ctc_m.sem_expected, 2)

check("angular_location 归尺寸类", _ctc_items[2].kind, C.KIND_DIM)
check("angular_location 走 dcr 2 跳",
      C._lookup_semantic_for_dev(_ctc_items[2], _ctc), ([118], "dim-2hop"))
check("短 ID 场景三率 100", (_ctc_m.recall, _ctc_m.precision), (100.0, 100.0))
check("短 ID 场景无断链", _ctc_paths.get("none", 0), 0)


# ============================================================
# 多视图复用 vs 同分组真重复
#   NIST ctc_02 里基准目标 A1~B4 同时挂在 DRAUGHTING_MODEL MBD_A 与 MBD_B 下
#   （SFA 的 ta 表 Saved Views 列写作 `(2) camera_model_d3 28 (MBD_A) 29 (MBD_B)`），
#   开发侧按「视图 × 标注」导出必然出两条 —— 这是合法结构，不能判成提取缺陷。
#   只有**同一分组内**同一 handle 出现多条才是真重复。
# ============================================================
_MV = """## MBD_A

- 标注数量：3 条

### 1. A1

detailData:
{
  "handle": "1167",
  "name": "Datum Target.1",
  "type": "datum_feature",
  "datum": "A",
  "target id": "1",
  "original_type_name": "draughting_callout"
}

### 2. Ø10 ±.1

detailData:
{
  "handle": "900",
  "name": "Linear Size.1",
  "type": "dimensional_size",
  "original_type_name": "draughting_callout"
}

### 3. Ø10 ±.1

detailData:
{
  "handle": "900",
  "name": "Linear Size.1",
  "type": "dimensional_size",
  "original_type_name": "draughting_callout"
}

## MBD_B

- 标注数量：1 条

### 1. A1

detailData:
{
  "handle": "1167",
  "name": "Datum Target.1",
  "type": "datum_feature",
  "datum": "A",
  "target id": "1",
  "original_type_name": "draughting_callout"
}
"""

_mvi, _mvd = C.parse_dev_markdown_ex(_MV)
check("多视图 解析到两个分组", _mvd.groups, ["MBD_A", "MBD_B"])
check("多视图 条目数", len(_mvi), 4)

_mvt = C.SfaTruth()
_mvt.semantic[500] = {"entity": "datum_target", "text": "A1 (point)",
                      "kind": C.KIND_DATUM}
_mvt.ta["Datum Target.1"] = {"id": 1028, "sem_refs": [500], "text": "A1"}
_mvt.dc[1167] = "Datum Target.1"

_mvr = C.match_items(_mvt, _mvi)
_mv_by_key = {}
for _r in _mvr:
    # 带语义命中的条目 key 会是 `MBD_A#1@500`（复合公差拆行），这里按开发条目归并
    _mv_by_key.setdefault(_r.key.split("@")[0], _r)

check("多视图 跨视图复用不判 DUP",
      [r.key for r in _mvr if C.DF_DUP in r.defects], ["MBD_A#2", "MBD_A#3"])
check("多视图 两种视图都标出复用了哪些分组",
      (_mv_by_key["MBD_A#1"].multiview, _mv_by_key["MBD_B#1"].multiview),
      ("MBD_A、MBD_B", "MBD_A、MBD_B"))
check("多视图 跨视图条目备注给出说明",
      "多个保存视图中复用" in _mv_by_key["MBD_B#1"].remark, True)
check("多视图 跨视图条目仍判命中",
      _mv_by_key["MBD_B#1"].status, C.ST_HIT)
check("多视图 同分组重复的明细指明分组",
      "同一分组（MBD_A）内出现 2 条" in _mv_by_key["MBD_A#2"].defect_detail, True)
check("多视图 同分组重复不算多视图", _mv_by_key["MBD_A#2"].multiview, "")
check("多视图 同分组重复不算多视图（第三条）", _mv_by_key["MBD_A#3"].multiview, "")

_mvm = C.compute_metrics(_mvr)
check("多视图 重复不计入缺陷率分母以外的额外项",
      _mvm.defect_by_code.get(C.DF_DUP), 2)
check("多视图 跨视图复用只算一条基准（不按视图数翻倍）", _mvm.datum_expected, 1)
check("多视图 基准覆盖率仍 100%", round(_mvm.datum_coverage, 2), 100.00)

# 回归：单分组、无重复时不得出现 DUP / 多视图标记
_solo = C.parse_dev_markdown(_MV.split("## MBD_B")[0])[:1]
_solo_r = C.match_items(_mvt, _solo)
check("多视图 单条标注不产生 DUP/多视图",
      [(r.key.split("@")[0], r.defects, r.multiview) for r in _solo_r],
      [("MBD_A#1", [], "")])

# 尺寸类（GT/DIM）走的是「同一语义实体被覆盖」分支，也要区分两种重复
_mv2 = C.SfaTruth()
_mv2.semantic[700] = {"entity": "dimensional_characteristic_representation",
                      "text": "⌀10 ±.1", "kind": C.KIND_DIM}
_mv2.ta["Linear Size.1"] = {"id": 900, "sem_refs": [901], "text": "⌀10 ±.1"}
_mv2.dcr_by_dim[901] = 700

_mv2r = C.match_items(_mv2, _mvi)
_mv2_remark = {r.key.split("@")[0]: r.remark for r in _mv2r}
_mv2_status = {r.key.split("@")[0]: r.status for r in _mv2r}
check("多视图 同分组重复：内容一致仍判命中，备注补 DUP 提示",
      _mv2_remark.get("MBD_A#3", ""),
      "ID 关联 + 语义指纹一致；同一 handle（900）重复导出（见缺陷 DUP）")
check("多视图 同分组重复不因「二次命中」被降级为文本差异",
      _mv2_status.get("MBD_A#3", ""), C.ST_HIT)

_MV3 = """## MBD_A

- 标注数量：1 条

### 1. Ø10 ±.1

detailData:
{
  "handle": "900",
  "name": "Linear Size.1",
  "type": "dimensional_size",
  "original_type_name": "draughting_callout"
}

## MBD_B

- 标注数量：1 条

### 1. Ø10 ±.1

detailData:
{
  "handle": "900",
  "name": "Linear Size.1",
  "type": "dimensional_size",
  "original_type_name": "draughting_callout"
}
"""
_mv3 = C.SfaTruth()
_mv3.semantic[700] = {"entity": "dimensional_characteristic_representation",
                      "text": "⌀10 ±.1", "kind": C.KIND_DIM}
_mv3.ta["Linear Size.1"] = {"id": 900, "sem_refs": [901], "text": "⌀10 ±.1"}
_mv3.dcr_by_dim[901] = 700
_mv3r = C.match_items(_mv3, C.parse_dev_markdown(_MV3))
_mv3_rem = {r.key.split("@")[0]: r.remark for r in _mv3r}
_mv3_status = {r.key.split("@")[0]: r.status for r in _mv3r}
check("多视图 跨视图复用且内容一致 → 判命中（不再无条件标差异）",
      _mv3_status.get("MBD_B#1", ""), C.ST_HIT)
check("多视图 跨视图的备注标明是复用而非重复导出",
      _mv3_rem.get("MBD_B#1", ""),
      "ID 关联 + 语义指纹一致 ；该标注在多个保存视图中复用（MBD_A、MBD_B），非重复导出")
check("多视图 跨视图的尺寸不判 DUP",
      [r.key for r in _mv3r if C.DF_DUP in r.defects], [])


# ============================================================
# 六、基准目标实体名漂移 / 标识两侧缀
#     覆盖 ctc_02 的整批失配：同一个实体在 SFA 报告里有两种写法 ——
#     ftc_10 写 `datum_target`，ctc_02 / ctc_05 / ftc_06 写
#     `placed_datum_target_feature`。旧代码只认前者，引用完全能落到语义表
#     （装载与引用校验全绿，体检面板看不出问题），但判据不认实体名，
#     结果开发侧全判「多余」、SFA 侧全判「缺失」。
#     同批还要处理标识两侧的干扰：`A1 (point)` 的目标形式后缀、
#     `⌀85\nK1` 与 `1.25x2\nC1` 的目标尺寸前缀 —— 按位置取首片段会拿到尺寸。
# ============================================================
def _build_sfa_dt(tmpdir: str) -> str:
    """基准目标实体名用 `placed_datum_target_feature`（ctc_02 的写法）。"""
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Semantic PMI Summary")
    ws.append(["Semantic PMI Summary (3)", None, None, None])
    ws.append(["ID", "Entity", "Semantic PMI", "Similar"])
    for rid, txt in [
        (3001, "A1 (point)"),          # 标识 + 目标形式
        (3002, "⌀85\nK1"),             # 目标尺寸 + 标识
        (3003, "1.25x2\nC1"),          # 目标尺寸（矩形） + 标识
    ]:
        ws.append([str(rid), "placed_datum_target_feature", txt, None])

    ws = wb.create_sheet("draughting_callout")
    ws.append(["draughting_callout  (3)", None])
    for rid, nm in [(3000, "Datum Target.1"), (3100, "Datum Target.2"),
                    (3200, "Datum Target.3")]:
        ws.append([str(rid), nm])

    ws = wb.create_sheet("tessellated_annotation_occurren")
    ws.append(["tessellated_annotation_occurrence (3)"] + [None] * 11)
    ws.append(["ID", "name", "styles", "item", "name", "children",
               "presentation style", "color", "plane", "Associated Geometry",
               "Associated Semantic PMI", "Saved Views"])
    for rid, nm, ref in [
        (3000, "Datum Target.1", "placed_datum_target_feature 3001"),
        (3100, "Datum Target.2", "placed_datum_target_feature 3002"),
        (3200, "Datum Target.3", "placed_datum_target_feature 3003"),
    ]:
        ws.append([str(rid), nm] + [""] * 8 + [ref, ""])

    path = os.path.join(tmpdir, "dt_sfa.xlsx")
    wb.save(path)
    return path


_MD_DT = """# PMI 提取结果

## MBD_X

- 标注数量：3 条

### 1. A1

detailData:
{
  "handle": "3000",
  "datum": "A",
  "name": "Datum Target.1",
  "target id": "1",
  "type": "datum_target"
}

### 2. K1

detailData:
{
  "handle": "3100",
  "datum": "K",
  "name": "Datum Target.2",
  "target id": "1",
  "type": "datum_target"
}

### 3. C1

detailData:
{
  "handle": "3200",
  "datum": "C",
  "name": "Datum Target.3",
  "target id": "1",
  "type": "datum_target"
}
"""

# 实体名判据：两种写法都要认；无关的 datum_feature / datum_system 不能误判
check("实体名 datum_target 认作基准目标",
      C.is_datum_target_entity("datum_target"), True)
check("实体名 placed_datum_target_feature 认作基准目标",
      C.is_datum_target_entity("placed_datum_target_feature"), True)
check("实体名 datum_feature 不误判",
      C.is_datum_target_entity("datum_feature"), False)
check("实体名 datum_system 不误判",
      C.is_datum_target_entity("datum_system"), False)
check("实体归类 基准目标归基准层",
      C.classify_sfa_entity("placed_datum_target_feature"), C.KIND_DATUM)
check("实体归类 未知名字不被认作已知",
      C.is_known_entity("some_future_entity"), False)
check("实体归类 已知名字通过",
      all(C.is_known_entity(x) for x in
          ("datum_system", "datum_target", "placed_datum_target_feature",
           "dimensional_characteristic_representation", "flatness_tolerance")), True)

# 标识提取：形式后缀与目标尺寸都不能顶替标识
check("标识提取 形式后缀 (point)", C.datum_target_ident("A1 (point)", "A1"), "A1")
check("标识提取 形式后缀 (area)", C.datum_target_ident("K1 (area)", "K1"), "K1")
check("标识提取 目标尺寸在标识前", C.datum_target_ident("D85 K1", "K1"), "K1")
check("标识提取 矩形目标尺寸在标识前",
      C.datum_target_ident("1.25x2 C1", "C1"), "C1")
check("标识提取 形式与尺寸同时出现",
      C.datum_target_ident("G1 (circular curve) (D = 1.)", "G1"), "G1")
check("标识提取 无期望值时按形态兜底",
      C.datum_target_ident("B4 (point)"), "B4")
check("标识提取 只有尺寸时给空",
      C.datum_target_ident("D85"), "")

with tempfile.TemporaryDirectory() as _td2:
    _truth_dt = C.load_sfa(_build_sfa_dt(_td2))

_items_dt, _ = C.parse_dev_markdown_ex(_MD_DT)
_rows_dt = C.match_items(_truth_dt, _items_dt)
_m_dt = C.compute_metrics(_rows_dt)
check("基准目标漂移 走 1 跳引用",
      C._lookup_semantic_for_dev(_items_dt[0], _truth_dt)[1], "datum_target-1hop")
check("基准目标漂移 三条全部命中",
      [r.status for r in _rows_dt if r.dev_name], [C.ST_HIT] * 3)
check("基准目标漂移 不产生反向缺失",
      [r.key for r in _rows_dt if r.status == C.ST_MISS], [])
check("基准目标漂移 基准覆盖率 100", _m_dt.datum_coverage, 100.0)
check("基准目标漂移 归基准层、不进语义层分母", _m_dt.sem_expected, 0)
check("基准目标漂移 SFA 侧无多余",
      [r.key for r in _rows_dt if r.status == C.ST_EXTRA], [])
check("基准目标漂移 体检项通过",
      [c.ok for c in _truth_dt.checks if c.name == "语义表实体归类"], [True])


# ============================================================
# 七、DMIA 关联通道（非 tessellated 导出）/ ◎ 与 ⭩◎ 的区分
#     覆盖 ctc_05-e1 的整批失配：该报告**没有** tessellated_annotation_occurrence
#     表，只有 draughting_model_item_association。旧代码只认前者 → t.ta 全空 →
#     开发侧全判「多余」、SFA 侧全判「缺失」，指标 0%，
#     而装载 / 引用 / name 三项交叉校验**全绿** —— 比「整表装载 0 条」更隐蔽。
#     DMIA 与 ta 同角色但结构不同：一条 callout 有多行（几何引用 + 语义引用），
#     必须按 callout 聚合，装载校验的分母也要跟着变成「分组合数」，
#     否则 6 行 → 3 组会被误判成「列定位错位」。
# ============================================================
def _build_sfa_dmia(tmpdir: str) -> str:
    """只给 DMIA、不给 tessellated 表的报告（ctc_05-e1 的形态）。"""
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Semantic PMI Summary")
    ws.append(["Semantic PMI Summary (3)", None, None, None])
    ws.append(["ID", "Entity", "Semantic PMI", "Similar"])
    for rid, ent, txt in [
        (4001, "position_tolerance", "⌖ | ⌀0.8 | A"),
        (4002, "placed_datum_target_feature", "1.25x2\nC1"),
        (4003, "dimensional_characteristic_representation", "⌀10.000 ± .001"),
    ]:
        ws.append([str(rid), ent, txt, None])

    ws = wb.create_sheet("draughting_callout")
    ws.append(["draughting_callout  (3)", None])
    ws.append(["ID", "name"])
    for rid, nm in [(4001, "Position.1"), (4002, "Datum Target C1 (15)"),
                    (4003, "Vertical Dimension (28)")]:
        ws.append([str(rid), nm])

    # 一条 callout 两行：一行给几何引用（shape_aspect），一行给语义引用
    ws = wb.create_sheet("draughting_model_item_associati")
    ws.append(["draughting_model_item_association  (6)"] + [None] * 5)
    ws.append(["ID", "name", "description", "definition", "used_representation",
               "identified_item"])
    for rid, tgt, dfn in [
        (501, 4001, "shape_aspect 5001"),
        (502, 4001, "position_tolerance 4001"),
        (503, 4002, "shape_aspect 5002"),
        (504, 4002, "placed_datum_target_feature 4002"),
        (505, 4003, "shape_aspect 5003"),
        (506, 4003, "dimensional_size 6003"),
    ]:
        ws.append([str(rid), "", "", dfn, "draughting_model 99",
                   f"draughting_callout {tgt}"])

    # shape_aspect 是 DMIA 必带的几何引用，须登记成「合法但无语义」的实体
    ws = wb.create_sheet("shape_aspect")
    ws.append(["shape_aspect  (3)", None])
    ws.append(["ID", "name"])
    for rid in (5001, 5002, 5003):
        ws.append([str(rid), ""])

    ws = wb.create_sheet("dimensional_characteristic_repr")
    ws.append(["dimensional_characteristic_representation  (1)"] + [None] * 2)
    ws.append(["ID", "dimension", "representation"])
    ws.append(["4003", "dimensional_size 6003", "shape_dimension_representation 7001"])

    path = os.path.join(tmpdir, "dmia_sfa.xlsx")
    wb.save(path)
    return path


def _build_sfa_nolink(tmpdir: str) -> str:
    """两种关联表都没有的报告：必须显式报断链，不能静默出 0% 结果。"""
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Semantic PMI Summary")
    ws.append(["Semantic PMI Summary (1)", None, None, None])
    ws.append(["ID", "Entity", "Semantic PMI", "Similar"])
    ws.append(["4001", "position_tolerance", "⌖ | ⌀0.8 | A", None])
    ws = wb.create_sheet("draughting_callout")
    ws.append(["draughting_callout  (1)", None])
    ws.append(["ID", "name"])
    ws.append(["4001", "Position.1"])
    path = os.path.join(tmpdir, "nolink_sfa.xlsx")
    wb.save(path)
    return path


_DMIA_MD = """# PMI 提取结果

## MBD_A

- 标注数量：3 条

### 1. ⌖ ⌀0.8 A

detailData:
{
  "handle": "4001",
  "name": "Position.1",
  "type": "position_tolerance"
}

### 2. C1

detailData:
{
  "handle": "4002",
  "datum": "C",
  "name": "Datum Target C1 (15)",
  "target id": "1",
  "type": "placed_datum_target_feature"
}

### 3. ⌀10.000 ±.001

detailData:
{
  "handle": "4003",
  "name": "Vertical Dimension (28)",
  "type": "dimensional_size"
}
"""

with tempfile.TemporaryDirectory() as _td_dm:
    _dmia = C.load_sfa(_build_sfa_dmia(_td_dm))
    _dmia_items, _ = C.parse_dev_markdown_ex(_DMIA_MD)
    _dmia_rows = C.match_items(_dmia, _dmia_items)
    _dmia_m = C.compute_metrics(_dmia_rows)
    _nolink = C.load_sfa(_build_sfa_nolink(_td_dm))

check("DMIA 通道被识别", _dmia.link_channel, C.LINK_DMIA)
check("DMIA ta 条数 = callout 数（按 callout 聚合，不是行数）", len(_dmia.ta), 3)
check("DMIA 一条 callout 的多行引用被合并",
      _dmia.ta["Position.1"]["sem_refs"], [5001, 4001])
check("DMIA 装载校验的分母是分组合数",
      [c.value for c in _dmia.checks if c.name.endswith("draughting_model_item_associati 装载")],
      ["3/3 条，聚合为 3 组"])
check("DMIA 交叉校验全通过", [b.line() for b in _dmia.broken_checks()], [])
check("DMIA 引用可解释（几何引用已登记为合法实体）",
      [c.value for c in _dmia.checks if c.name == "ta 引用可解释"], ["6/6 = 100%"])
check("DMIA 三条开发条目全部命中",
      [r.status for r in _dmia_rows if r.dev_name], [C.ST_HIT] * 3)
check("DMIA 关联路径齐全",
      sorted({r.path for r in _dmia_rows}),
      ["datum_target-1hop", "dim-2hop", "gt-1hop"])
check("DMIA 三项指标 100",
      (round(_dmia_m.recall, 2), round(_dmia_m.precision, 2),
       round(_dmia_m.datum_coverage, 2)), (100.0, 100.0, 100.0))

check("无关联通道时体检报错而非静默",
      [c.name for c in _nolink.broken_checks()], ["图形标注关联通道"])
check("无关联通道时给出明确告警",
      any("两种图形标注关联表都没有" in w for w in _nolink.warnings), True)

# 同心度符号：`⭩◎` 是复合公差的渲染产物，其中的 ◎ 与同心度无关；
# 独立出现的 ◎（ctc_05 `◎ | ⌀0.03 | A`）才是真符号。
check("◎ 独立出现算同心度符号",
      sorted(C.symbols_of("◎ | ⌀0.03 | A")), ["⌀", "◎"])
check("⭩◎ 渲染产物不算符号",
      sorted(C.symbols_of("⭩◎ | ⌓ | 1.5 | A")), ["⌓"])
# ASME 的 Unicode 形位公差写法：圆跳动 ↗、全跳动 ⌰、圆度 ○
check("新增 ASME 符号被识别（圆跳动/全跳动/圆度）",
      sorted(C.symbols_of("↗ | 0.035 | A-B") | C.symbols_of("⌰ | 0.015 | B")
             | C.symbols_of("○ | 0.002")),
      ["↗", "⌰", "○"])


# ============================================================
# 八、图形专用导出（ta 表在、但没有关联键列）
#     覆盖 ftc_08 `-e1-tg`：ta 表 12 列，表头里**没有** `Associated Semantic PMI`，
#     整份报告也没有 `Semantic PMI Summary`。三个坑叠在一起：
#       1. `_locate` 过去在表头识别成功时仍给缺列套默认列号 —— 关联键的默认值 10
#          正好落在 `Saved Views` 上，于是从 `camera_model_d3 58866 (MBD_A)` 抠出
#          一堆假引用（实测该报告 116 条），把「关联键不存在」伪装成「引用无法解释」；
#       2. SFA 侧零真值，开发侧每一条都被判「⚠️ 多余」，精确率 0% —— 看起来像
#          「开发侧全错提取」，实际是「选错了报告文件」；
#       3. `draughting_model_item_association` 指向的是 tessellated 标注本身而非
#          draughting_callout，兜不住底。
#     期望：不产生假引用、判「⛔ 无可比真值」、体检点名、并指出同目录里含语义 PMI
#     的那份报告。
# ============================================================
def _build_sfa_graphic_only(tmpdir: str,
                            name: str = "nist_ftc_08_asme1_ap242-e1-tg-sfa.xlsx") -> str:
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("tessellated_annotation_occurren")
    ws.append(["tessellated_annotation_occurrence  (2)"] + [None] * 11)
    ws.append([None] * 12)
    ws.append([f"列{i}" for i in range(1, 13)])
    # 表头照抄 ftc_08：**没有** Associated Semantic PMI / Equivalent Unicode String
    ws.append(["ID", "name", "styles", "item", "name", "children",
               "presentation style", "color", "plane", "Associated Geometry",
               "Saved Views", "Validation Properties"])
    ws.append(["25921", "Flatness.1", "(1) presentation_style_assignment 25916",
               "tessellated_geometric_set 25911", "flatness",
               "(63) tessellated_curve_set", "curve_style 25917",
               "draughting_pre_defined_colour 25919  (white)",
               "annotation_plane 27316\n(A1)",
               "(1) shape_aspect 55861\n(1) complex_triangulated_face 7581",
               "camera_model_d3 58866 (MBD_A)",
               "property definition 59376\n(pmi validation property)"])
    ws.append(["26231", "Simple Datum.1", "(1) presentation_style_assignment 26226",
               "tessellated_geometric_set 26221", "datum",
               "(10) tessellated_curve_set", "curve_style 26227",
               "draughting_pre_defined_colour 26229  (white)",
               "annotation_plane 27316\n(A1)",
               "(1) shape_aspect 55861",
               "camera_model_d3 58866 (MBD_A)",
               "property definition 59466\n(pmi validation property)"])

    # 关联表在，但指向 tessellated 标注本身 + shape_aspect（纯几何），兜不住语义
    ws = wb.create_sheet("draughting_model_item_associati")
    ws.append(["draughting_model_item_association  (2)"] + [None] * 5)
    ws.append(["ID", "name", "description", "definition", "used_representation",
               "identified_item"])
    for rid, nm, ta_id in [(55876, "Flatness.1", 25921),
                           (55886, "Simple Datum.1", 26231)]:
        ws.append([str(rid), nm, "", "shape_aspect 55861",
                   "draughting_model 25280",
                   f"tessellated_annotation_occurrence {ta_id}"])

    ws = wb.create_sheet("shape_aspect")
    ws.append(["shape_aspect  (2)"] + [None] * 4)
    ws.append(["ID", "name", "description", "of_shape", "product_definitional"])
    ws.append(["55861", "Flatness.1", "", "product_definition_shape 56", "True"])

    path = os.path.join(tmpdir, name)
    wb.save(path)
    return path


def _write_min_semantic(path: str) -> str:
    """最小「含语义 PMI」报告，只为验证「兄弟报告提示」能认出它。"""
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Semantic PMI Summary")
    ws.append(["Semantic PMI Summary (1)", None, None, None])
    ws.append(["ID", "Entity", "Semantic PMI", "Similar"])
    ws.append(["55861", "flatness_tolerance", "▱ | 0.03", None])
    wb.save(path)
    return path


# 列定位：表头在手时，找不到的列必须是 None，不能套默认列号
_cr_gap = C._locate([["ID", "name", "Saved Views"], ["1", "a", "x"]],
                    ("id", "name"),
                    ("id", "name", "associated semantic pmi"),
                    (0, 1, 10))
check("列名定位 缺列不再套默认列号",
      _cr_gap.resolved["associated semantic pmi"], None)
check("列名定位 缺列有留痕", _cr_gap.notes, ["表头未见 `associated semantic pmi`，"
      "该列按**不存在**处理（不套默认列号 10 —— 位置已变，按位置取会读到别的列）"])
check("列名定位 id 仍回落第 1 列",
      C._locate([["name", "item"], ["a", "b"]], ("id", "name"), ("id", "name"),
                (0, 1)).resolved["id"], 0)

with tempfile.TemporaryDirectory() as _td_go:
    _go_path = _build_sfa_graphic_only(_td_go)
    _sib_path = _write_min_semantic(
        os.path.join(_td_go, "nist_ftc_08_asme1_ap242-e2-sfa.xlsx"))
    _go = C.load_sfa(_go_path)
    _sug = C.suggest_reports_with_truth(_go_path)

check("图形专用导出 被识别为 graphic_only", _go.graphic_only, True)
check("图形专用导出 关联键列记为不存在",
      _go.recipe("ta").resolved["associated semantic pmi"], None)
check("图形专用导出 图形文本列记为不存在",
      _go.recipe("ta").resolved["equivalent unicode string"], None)
check("图形专用导出 不按默认列号从 Saved Views 抠假引用",
      [r["sem_refs"] for r in _go.ta.values()], [[], []])
check("图形专用导出 表内条目仍装载（有结构没内容）", _go.recipe("ta").n_loaded, 2)
check("图形专用导出 数据行仍能数出来", _go.recipe("ta").n_rows, 2)
check("图形专用导出 关联通道判定为不可用", _go.link_channel, "")
check("图形专用导出 判定为无真值", _go.truth_missing, True)
check("图形专用导出 体检点名四项",
      sorted(c.name for c in _go.broken_checks()),
      ["draughting_model_item_associati 装载", "ta 引用可解释",
       "图形标注关联通道", "图形标注关联键列", "语义真值"])
check("图形专用导出 装载告警不把「按设计产不出」说成列定位错位",
      [c.detail for c in _go.broken_checks()
       if c.name.endswith("draughting_model_item_associati 装载")],
      ["该报告没有 `draughting_callout` 表，关联表的 `identified_item` "
       "指向的是 tessellated 标注本身、而不是 callout，"
       "本表按设计产出不了关联条目（不是列定位错位）"])
check("图形专用导出 告警点名缺列",
      any("Associated Semantic PMI" in w for w in _go.warnings), True)
check("图形专用导出 告警挑明「不含语义 PMI」",
      any("不含任何语义 PMI" in w for w in _go.warnings), True)
check("导错文件时指出同目录含语义 PMI 的兄弟报告",
      [os.path.basename(x) for x in _sug],
      ["nist_ftc_08_asme1_ap242-e2-sfa.xlsx"])
check("兄弟报告提示写进了装载告警",
      any("nist_ftc_08_asme1_ap242-e2-sfa.xlsx" in w for w in _go.warnings), True)
check("兄弟报告提示不误报当前文件",
      all(os.path.abspath(x) != os.path.abspath(_go_path) for x in _sug), True)

# 报告与标准件目录不在一起时（用户常把导出件拖到桌面），回落到 PMI_SFA_DIR
with tempfile.TemporaryDirectory() as _td_a, tempfile.TemporaryDirectory() as _td_b:
    _go_other = _build_sfa_graphic_only(_td_a)
    _write_min_semantic(os.path.join(_td_b, "nist_ftc_08_asme1_ap242-e2-sfa.xlsx"))
    _env_old = os.environ.get("PMI_SFA_DIR")
    os.environ["PMI_SFA_DIR"] = _td_b
    try:
        _sug2 = C.suggest_reports_with_truth(_go_other)
    finally:
        if _env_old is None:
            os.environ.pop("PMI_SFA_DIR", None)
        else:
            os.environ["PMI_SFA_DIR"] = _env_old
check("同目录无候选时回落到 PMI_SFA_DIR",
      [os.path.basename(x) for x in _sug2], ["nist_ftc_08_asme1_ap242-e2-sfa.xlsx"])

_go_md = """# PMI 提取结果

## MBD_A

- 标注数量：3 条

### 1. ⏥ .03

detailData:
{
  "handle": "1233",
  "name": "Feature Control Frame (1)",
  "type": "flatness_tolerance"
}

### 2. B

detailData:
{
  "handle": "1235",
  "datum": "B",
  "name": "Datum Feature Symbol B (54)",
  "type": "datum_feature"
}

### 3. Note (55)

detailData:
{
  "handle": "1238",
  "description": "NOTES (UNLESS OTHERWISE SPECIFIED):",
  "name": "Note (55)",
  "property": "semantic text"
}
"""
_go_items, _ = C.parse_dev_markdown_ex(_go_md)
_go_rows = C.match_items(_go, _go_items)
_go_m = C.compute_metrics(_go_rows)

check("无真值 开发侧判「⛔ 无可比真值」而非「⚠️ 多余」",
      [r.status for r in _go_rows if r.kind != C.KIND_NOTE],
      [C.ST_ABSENT] * 2)
check("无真值 不产生任何「多余」", _go_m.sem_extra, 0)
check("无真值 不进语义指标分母", _go_m.sem_expected, 0)
check("无真值 计入不可比条目", _go_m.no_truth, 2)
check("无真值 指标标记为不成立", _go_m.truth_missing, True)
check("无真值 关联路径记为 no-truth",
      C.link_stats(_go_rows), {"note": 1, "no-truth": 2})
check("无真值 备注写明不代表提取错误",
      all("不代表开发侧提取错误" in r.remark for r in _go_rows
          if r.status == C.ST_ABSENT), True)
check("无真值 注释类仍按注释处理",
      [r.status for r in _go_rows if r.kind == C.KIND_NOTE], [C.ST_NOTE])
check("无真值 汇总标出指标不成立",
      C.summarize(_go_rows, _go, _go_items)["指标是否成立"],
      "否（SFA 报告不含语义 PMI）")

# 数量标注位置：SFA 放串首、开发侧可能渲染到串尾，不应判「数量前缀 开发=无」
_go_codes, _ = C.dev_defects(
    C.DevItem(title="⌀0.237 +.005 -0.001 2X", name="x",
              detail={"type": "dimensional_size"}),
    "2X ⌀0.237  +0.005 -0.001")
check("数量标注在串尾也能对上（不报 CNT）", _go_codes, [])
_go_codes2, _ = C.dev_defects(
    C.DevItem(title="⌀0.237 +.005 -0.001", name="x",
              detail={"type": "dimensional_size"}),
    "2X ⌀0.237  +0.005 -0.001")
check("一侧真的缺数量标注仍要报 CNT", _go_codes2, [C.DF_CNT])


# ============================================================
# 文件对象入参（UI 上传场景）
#
# Streamlit `file_uploader` 给的是 `UploadedFile` —— 一个 BytesIO 加 `.name`
# （上传时的原始文件名），**不是路径**。早期 `load_sfa` 上来就 `os.path.basename(path)`，
# UI 里一点「开始比对」就报 `SFA 报告解析失败：expected str, bytes or os.PathLike
# object, not UploadedFile`。内核必须同时吃路径和文件对象：
# 差别只在「同目录找兄弟报告」这类路径语义 —— 文件对象没有目录，那一路跳过。
# ============================================================
class _FakeUpload(io.BytesIO):
    """最小复刻 Streamlit UploadedFile：BytesIO + 上传时的原始文件名。"""

    def __init__(self, data: bytes, name: str):
        super().__init__(data)
        self.name = name


with tempfile.TemporaryDirectory() as _td_u:
    _up_path = _build_sfa(_td_u)
    with open(_up_path, "rb") as _f:
        _up_bytes = _f.read()
    _up_name = os.path.basename(_up_path)
    try:
        _t_file = C.load_sfa(_FakeUpload(_up_bytes, _up_name))
        _up_err = ""
    except Exception as _e:                    # noqa: BLE001
        _t_file, _up_err = None, f"{type(_e).__name__}: {_e}"
    _t_pathv = C.load_sfa(_up_path)

check("文件对象入参不再抛异常（UploadedFile 场景）", _up_err, "")
if _t_file is not None:
    check("文件对象入参 与路径入参 结果一致",
          (len(_t_file.semantic), len(_t_file.dc), len(_t_file.ta), len(_t_file.datum)),
          (len(_t_pathv.semantic), len(_t_pathv.dc), len(_t_pathv.ta),
           len(_t_pathv.datum)))
    check("文件对象入参 体检结论与路径入参一致",
          [c.name for c in _t_file.broken_checks() if not c.ok],
          [c.name for c in _t_pathv.broken_checks() if not c.ok])
    check("文件对象入参 保留上传时的文件名（供导错文件提示）",
          _t_file.path, _up_name)

check("_source_name 同时吃路径与文件对象",
      (C._source_name(r"C:\tmp\nist_ftc_08-sfa.xlsx"),
       C._source_name(_FakeUpload(b"", "nist_ftc_08-sfa.xlsx"))),
      ("nist_ftc_08-sfa.xlsx", "nist_ftc_08-sfa.xlsx"))

# 文件对象没有所在目录，「同目录找兄弟报告」必然落空 → 必须靠 PMI_SFA_DIR 兜住。
# 这正是 UI 的常态：导出件从浏览器传上来，标准件目录只能靠环境变量告诉工具。
with tempfile.TemporaryDirectory() as _td_u2:
    _write_min_semantic(os.path.join(_td_u2, "nist_ftc_08_asme1_ap242-e2-sfa.xlsx"))
    _env_old3 = os.environ.get("PMI_SFA_DIR")
    os.environ["PMI_SFA_DIR"] = _td_u2
    try:
        _sug3 = C.suggest_reports_with_truth(
            _FakeUpload(b"", "nist_ftc_08_asme1_ap242-e1-tg-sfa.xlsx"))
    finally:
        if _env_old3 is None:
            os.environ.pop("PMI_SFA_DIR", None)
        else:
            os.environ["PMI_SFA_DIR"] = _env_old3
check("文件对象入参：兄弟报告提示退化为扫 PMI_SFA_DIR",
      [os.path.basename(x) for x in _sug3], ["nist_ftc_08_asme1_ap242-e2-sfa.xlsx"])


# ============================================================
# 九、SFA 溯源标注 `(composite with <id>)`
#
# 复合公差的后续段，SFA 会在单元格末行标注它归属哪条复合公差：
#     ⌓ | 0.2 | A / ▽ / ⎹ / [D] / (composite with 486)
# 这是元数据不是 PMI 内容，且括号里带数字。旧版独立引擎（现已废弃的
# `pmi_compare_work.py` 上半部）里有这条清洗，迁到 pmi_core 时漏了，
# 于是「SFA 有、开发侧未提取」的反查行把它原样摊在结果表里。
# ============================================================
_COMP_RAW = "⌓ | 0.2 | A \n   ▽\n   ⎹\n   [D]\n(composite with 486)"

check("溯源标注 归一化后不再含 composite with",
      "composite" in C.normalize(_COMP_RAW), False)
check("溯源标注 归一化后不含括号内编号 486",
      "486" in C.normalize(_COMP_RAW), False)
check("溯源标注 归一化保留 FCF 本体与后段基准",
      C.normalize(_COMP_RAW), "⌓ | 0.2 | A ▽ ⎹ [D]")
check("溯源标注 指纹不把编号当公差数值",
      C.sem_tokens(_COMP_RAW)["numbers"], ["0.2"])
check("溯源标注 大小写 / 空格变体同样抹掉",
      C.strip_composite_marker("( COMPOSITE  with  488 )").strip(), "")
check("溯源标注 不误伤 [C] / <ST> 这类修饰符",
      C.strip_composite_marker("⌖ | ⌀0.02 Ⓜ | D | B | C [C] <ST>"),
      "⌖ | ⌀0.02 Ⓜ | D | B | C [C] <ST>")
check("溯源标注 sfa_text_for 取片段时也已抹掉",
      C.sfa_text_for(_COMP_RAW, C.KIND_GT), "⌓ | 0.2 | A")
# 单行写法（图形 ta 表可能整条渲染在一行）也必须干净
check("溯源标注 与 FCF 同行时取片段不含标注",
      C.sfa_text_for("⌓ | 0.12 | A (composite with 488)", C.KIND_GT),
      "⌓ | 0.12 | A")
# 反查行（SFA 有、开发侧无）走的是 normalize，展示层必须干净
check("溯源标注 反查行展示文本已清洗",
      "composite" in C.normalize("⌓ | 1.2 | A \n ▽\n ⎹\n [D]"), False)


# ============================================================
# 汇总
# ============================================================
for f in FAIL:
    print("FAIL -", f)
print(f"\n通过 {len(PASS)} / {len(PASS) + len(FAIL)}")
sys.exit(1 if FAIL else 0)

